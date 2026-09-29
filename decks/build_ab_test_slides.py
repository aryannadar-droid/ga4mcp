"""Five persona A/B test slides for the MREI CRO deck (13.33 x 7.5, Sora, deck palette)."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

OUT = sys.argv[1]
NAVY, RED, GREY, LIGHT = "0F2C5C", "C8102E", "5A6472", "96A1B0"
PANEL, TINT, LINE, WHITE = "F4F6FA", "E8EEF7", "D8DFE8", "FFFFFF"
GREEN, GREEN_T, AMBER, AMBER_T = "1E7B34", "E5F1EA", "B26B00", "FFF3DC"
FONT = "Sora"
warn = []

p = Presentation()
p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
BLANK = p.slide_layouts[6]

def rgb(h): return RGBColor.from_string(h)

def fits(t, pt, w, h):
    cpl = max(1, int((w * 72) / (pt * 0.56)))
    return sum(max(1, -(-len(l) // cpl)) for l in t.split("\n")) * pt * 1.28 / 72 <= h

def box(s, x, y, w, h, fill=None, line=None, r=0.06, dash=False):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.adjustments[0] = r
    if fill: b.fill.solid(); b.fill.fore_color.rgb = rgb(fill)
    else: b.fill.background()
    if line:
        b.line.color.rgb = rgb(line); b.line.width = Pt(1)
        if dash: b.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else: b.line.fill.background()
    b.shadow.inherit = False
    return b

def text(s, x, y, w, h, paras, size=9, color=NAVY, bold=False, align="l", anchor="t", italic=False, sp=0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[anchor]
    paras = paras if isinstance(paras, list) else [paras]
    plain = []
    for i, para in enumerate(paras):
        pp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pp.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}[align]
        if sp: pp.space_after = Pt(sp)
        segs = para if isinstance(para, list) else [(para, {})]
        ln = ""
        for seg, o in segs:
            r = pp.add_run(); r.text = seg; ln += seg
            f = r.font; f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic); f.color.rgb = rgb(o.get("color", color))
        plain.append(ln)
    if not fits("\n".join(plain), size, w, h + 0.02):
        warn.append(f"fit? {plain[0][:45]!r} {size}pt {w}x{h}")
    return tb

CHIP = {"MEASURED": (GREEN, GREEN_T), "HYPOTHESIS": (NAVY, TINT), "NEEDS MRIIRS DATA": (AMBER, AMBER_T),
        "SHIP-READY": (GREEN, GREEN_T), "WAVE 1": (WHITE, RED), "WAVE 2": (WHITE, NAVY)}
def chip(s, label, x, y, h=0.22):
    fg, bg = CHIP[label]
    w = 0.22 + 0.072 * len(label)
    b = box(s, x, y, w, h, fill=bg, line=(None if bg in (RED, NAVY) else fg), r=0.5)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    r = pp.add_run(); r.text = label
    r.font.name = FONT; r.font.size = Pt(6.5); r.font.bold = True; r.font.color.rgb = rgb(fg)
    return w

def test_column(s, x, t):
    W = 6.0
    box(s, x, 1.62, W, 5.30, fill=PANEL, r=0.03)
    # header
    b = box(s, x + 0.2, 1.80, 0.62, 0.36, fill=RED if t["wave"] == 1 else NAVY, r=0.25)
    text(s, x + 0.2, 1.80, 0.62, 0.36, t["id"], size=10, color=WHITE, bold=True, align="c", anchor="m")
    text(s, x + 0.95, 1.74, 2.9, 0.46, t["name"], size=11.5, color=NAVY, bold=True, anchor="m")
    cx = x + 0.2
    for c in t["chips"]:
        cx += chip(s, c, cx, 2.34) + 0.08
    text(s, x + 0.2, 2.64, 3.62, 0.2, t["from"], size=7, color=LIGHT, italic=True)
    # sections
    y = 2.98
    TW = 3.62
    for label, body in [("WHAT WE SEE", t["see"]), ("CONTROL → VARIANT", t["change"]),
                        ("HYPOTHESIS", t["hyp"]), ("MEASURE", t["measure"]), ("SET-UP", t["setup"])]:
        cpl = int(TW * 72 / (8 * 0.54))
        h = (-(-len(body) // cpl)) * 8 * 1.3 / 72
        text(s, x + 0.2, y, TW, 0.16, label, size=7, color=RED, bold=True)
        text(s, x + 0.2, y + 0.17, TW, h, body, size=8, color=(NAVY if label != "WHAT WE SEE" else GREY),
             italic=(label == "HYPOTHESIS"))
        y += 0.17 + h + 0.12
    if y > 6.95: warn.append(f"{t['id']} column overflows to {y:.2f}")
    # mockup frame (phone ratio 390x844)
    mx, my, mw, mh = x + 3.98, 1.80, 1.84, 3.98
    box(s, mx, my, mw, mh, fill=WHITE, line=LIGHT, r=0.1, dash=True)
    text(s, mx + 0.1, my + mh / 2 - 0.35, mw - 0.2, 0.7,
         [[("MOCKUP", {"bold": True, "color": LIGHT, "size": 8})], [(t["mock"], {})]],
         size=7, color=LIGHT, align="c")
    text(s, mx, my + mh + 0.08, mw, 0.95, t["page"], size=7, color=GREY, align="c")

def slide(persona, question, sub, tests, note, source, notes):
    s = p.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(WHITE)
    text(s, 0.60, 0.32, 9.0, 0.20, f"EXPERIMENTATION  ·  A/B TESTS  ·  {persona.upper()}", size=9, color=RED, bold=True)
    chip(s, "HYPOTHESIS", 12.73 - (0.22 + 0.072 * 10), 0.30)
    text(s, 0.60, 0.55, 12.10, 0.60, question, size=24, color=NAVY, bold=True)
    text(s, 0.60, 1.17, 12.10, 0.30, sub, size=12, color=GREY)
    test_column(s, 0.60, tests[0])
    test_column(s, 6.73, tests[1])
    text(s, 0.60, 7.00, 12.10, 0.30,
         [[("ORDER  ", {"bold": True, "color": NAVY}), (note + "   ", {}), ("Source  ", {"bold": True, "color": NAVY}), (source, {})]],
         size=7, color=LIGHT)
    s.notes_slide.notes_text_frame.text = notes

RULE = ("Runs after Phase 0 (generate_lead, unified phone_click, brochure_request live; Fee Calculator redirect fixed). "
        "Thresholds are proposals to agree with MREI before launch: scale at 95% confidence with the guardrail within −5% relative.")

# ---------------------------------------------------------------- 1 Course Researcher
slide("Course Researcher",
      "Course Researcher: finish the decision on the course page",
      "“What is this course, and is it worth it?” · B.Tech CSE · 12,308 landings · 53,812 persona sessions",
      [dict(id="CR-1", wave=1, name="Fee summary at decision height", chips=["WAVE 1", "SHIP-READY"],
            from_="", **{"from": "From MREI T5 · pitch Fix 04 (whole cost)"},
            see="84.3% of onward moves go to the fee page; 99.3% come straight back; 6 continue. The fee block sits at 11,288px of 16,456px (~13 screens) and shows year 1 only.",
            change="“Fee Details” anchor to the bottom block → fold-2 fee card: Year 1 ₹3,05,800 · instalments ₹1,85,900 + ₹1,19,900 · industry-partner specialisation ₹3,33,800. Beside it: “Send this fee sheet on WhatsApp” and “Check eligibility”.",
            hyp="If the full CSE fee sits where the decision is made, leads on the CSE page rise, because users stop leaving for the fee table and the journey no longer ends on the return trip.",
            measure="Primary: leads on the CSE page · Secondary: share of onward moves to the fee page (84.3%) · Guardrail: engagement (67.9%) · Scale if leads ≥ +10% relative",
            setup="50/50, all CSE traffic, mobile + desktop, 4 weeks. Live table values only; add years 2–4 and hostel when MREI supplies them.",
            h=(0.52, 0.78, 0.52, 0.52, 0.40),
            mock="Control: fold 2 today\nVariant: inline fee card + WhatsApp / eligibility CTAs",
            page="/mriirs/academics/btech-computer-science-and-engineering · 390px"),
       dict(id="CR-2", wave=2, name="Brochure + optional WhatsApp send, no gate", chips=["WAVE 2", "SHIP-READY"],
            **{"from": "From MREI T6 · pitch N08 (reciprocity)"},
            see="483 brochure downloads vs 47 Apply Now clicks on this page. Course pages: 516 file downloads vs 64 form starts. The brochure is an anonymous PDF today.",
            change="Direct PDF link → PDF opens at once, then a bottom sheet: “Get the CSE brochure + fee sheet on WhatsApp”. Mobile number only, consent box unticked, course pre-filled.",
            hyp="If downloaders are offered the brochure on WhatsApp without a gate, known leads rise, because the strongest intent signal on the page finally records who acted.",
            measure="Primary: WhatsApp opt-ins → leads · Secondary: brochure_request · Guardrail: brochure downloads must not fall > 5% · Keep if incremental leads ≥ +8%",
            setup="50/50 on brochure clicks, 3 weeks, after CR-1 decides. Effort S: WhatsApp template + one bottom sheet.",
            h=(0.52, 0.66, 0.52, 0.52, 0.40),
            mock="Control: tap → PDF\nVariant: PDF + WhatsApp bottom sheet",
            page="Brochure (PDF) link, fold 1 of the CSE page")],
      "CR-1 first; CR-2 once CR-1 decides (same template).",
      "GA4 explorations 13 Jun – 10 Sept 2026 · GA4 event export by page type · live audit 24 Sept 2026 · MRIIRS pitch v4.",
      "Course Researcher tests. CR-1 answers the round-trip to the fee page; CR-2 captures the download moment without gating information. " + RULE)

# ---------------------------------------------------------------- 2 Cost Checker
slide("Cost Checker",
      "Cost Checker: one fee, found fast, with a person beside it",
      "“What does it cost, and can I get in?” · fee-structure pages · 8,664 + 3,752 sessions · 74% engagement",
      [dict(id="CC-1", wave=1, name="Fee finder instead of the 15-column table", chips=["WAVE 1", "SHIP-READY"],
            **{"from": "From MREI T1 + T4 · pitch T3 (hostel toggle)"},
            see="70% of MRIIRS fee entries come from CSE. Search “CSE” returns unrelated rows; the total is column 15 of a 1,642px table. 1,236 moves go to the homepage.",
            change="Table + free-text search → Institution → programme → specialisation selectors → one fee card: total first, instalments, hostel toggle updates the total, “← Back to B.Tech CSE”. Pre-filled via ?programme=btech-cse from the CSE page.",
            hyp="If a Cost Checker reaches their programme’s fee in three taps, leads rise and homepage exits fall, because the number they came for is no longer hidden.",
            measure="Primary: leads on fee pages · Secondary: fee → homepage exits (1,236) · Guardrail: engagement (73.9%) · Scale if leads ≥ +10% relative",
            setup="50/50 on /mriirs and /mru fee pages, 4 weeks. Needs the fee_search event and the calculator redirect fixed first.",
            h=(0.52, 0.78, 0.52, 0.52, 0.40),
            mock="Control: table scrolled\nVariant: selectors + one fee card",
            page="/mriirs/admissions/fee-structure · /mru/admissions/fee-structure"),
       dict(id="CC-2", wave=2, name="Call and callback under the fee", chips=["WAVE 2", "SHIP-READY"],
            **{"from": "From MREI T2 + T3 · pitch Fix 09 (hand-off)"},
            see="82 phone clicks vs 31 form submits: 12.5% of all phone clicks from 6.7% of landings — the only page type where calls beat forms. The number is in the footer only.",
            change="Footer number + generic Apply / Enquire tabs → “Questions about B.Tech CSE fees?” block under the fee card: tap-to-call (tracked with page type) · “Call me back in 10 min” (name + mobile, programme pre-filled) · “Am I eligible?”.",
            hyp="If the call route sits beside the fee, contact rises, because at the money moment users want a person, not a form.",
            measure="Primary: leads (callbacks + submits) · Secondary: phone clicks · Guardrail: total contact actions · Ops: callback SLA agreed with admissions",
            setup="Mobile only, 50/50, 3 weeks, after CC-1 decides. Effort S: one block + callback routing in NoPaperForms.",
            h=(0.52, 0.78, 0.40, 0.52, 0.40),
            mock="Control: fee rows + footer number\nVariant: call / callback / eligibility block",
            page="Fee-structure pages, mobile, below the fee card")],
      "CC-1 first; CC-2 on the same template once CC-1 decides.",
      "GA4 explorations 13 Jun – 10 Sept 2026 · GA4 event export by page type · live audit 24 Sept 2026 · MRIIRS pitch v4.",
      "Cost Checker tests. The highest-intent visitor on the site. CC-1 fixes findability; CC-2 gives them the person they already reach for. " + RULE)

# ---------------------------------------------------------------- 3 Early Explorer
slide("Early Explorer",
      "Early Explorer: every answer should lead to a programme",
      "“What should I even study after 12th?” · blog article template · 792 articles · 22,813 persona sessions",
      [dict(id="EE-1", wave=1, name="Programme card inside the article", chips=["WAVE 1", "SHIP-READY"],
            **{"from": "From MREI T10 · Amity “Highlights” pattern"},
            see="35 of 4,195 NEET/BPT readers reach the BPT page. Articles carry ~10 blog links to 0–1 programme links. The BDS article names “MBA in Hospital Management” with no link.",
            change="“Also Read” blog links → Highlights card after the section that answers the question, repeated at the end: programme · duration · eligibility · year-1 fee · “Explore BPT at MRIIRS →”. Top 20 articles by sessions, from an article → programme mapping sheet.",
            hyp="If the article names the matching programme where the question is answered, article → programme progression rises, because readers get a next step other than the homepage.",
            measure="Primary: article → programme rate (0.8% today) · Secondary: form_start · Guardrail: engagement · Scale if ≥ +50% relative",
            setup="50/50 on the top 20 articles, 4 weeks. Precondition: validate blog traffic first (32,318 index landings, 42 scroll events).",
            h=(0.52, 0.78, 0.52, 0.40, 0.40),
            mock="Control: article + Also Read\nVariant: mid + end programme card",
            page="Blog article, e.g. “Is NEET Compulsory for BPT Admission 2026?”"),
       dict(id="EE-2", wave=2, name="“Help me choose” instead of Apply Now", chips=["WAVE 2", "SHIP-READY"],
            **{"from": "From MREI T11 · pitch T8 (track quiz) · Amity 1-1 Counselling"},
            see="Blog landings produced 0 Apply and 0 phone clicks. The sticky CTA asks an undecided reader to apply. Amity swaps it for “1-1 Counselling” on this template.",
            change="Sticky Apply Now / Enquire Now → “Not sure what to study? Take the 5-question course guide”. Result: 3 matching MRIIRS / MRU programmes, sent on WhatsApp.",
            hyp="If undecided readers are offered a guide instead of Apply, lead capture rises, because the ask matches where they are in the decision.",
            measure="Primary: quiz completion → lead · Secondary: lead → application, quiz leads vs form leads · Keep if completion ≥ 40% and quiz leads convert ≥ form leads",
            setup="50/50 on the blog template, 4 weeks, after EE-1. Effort M: quiz logic + programme mapping + WhatsApp result.",
            h=(0.52, 0.52, 0.52, 0.52, 0.40),
            mock="Control: sticky Apply / Enquire\nVariant: sticky course guide + quiz step",
            page="Blog template, sticky bottom bar")],
      "Blog traffic check first, then EE-1, then EE-2.",
      "GA4 explorations 13 Jun – 10 Sept 2026 · GA4 event export by page type · live audit 24 Sept 2026 · Amity Online blog, 22 Sept 2026 · MRIIRS pitch v4.",
      "Early Explorer tests. The blog answers its own question and then sends readers home. EE-1 adds the bridge; EE-2 changes the ask. " + RULE)

# ---------------------------------------------------------------- 4 Proof Seeker
slide("Proof Seeker",
      "Proof Seeker: attach verified proof to the programme",
      "“Should I believe what they claim?” · awards page and B.Tech CSE · 27,579 persona sessions",
      [dict(id="PS-1", wave=1, name="Link each award to its programmes", chips=["WAVE 1", "SHIP-READY"],
            **{"from": "From MREI T9 · Amity text badges"},
            see="/mriirs/about/awards-recognitions: 1,500 sessions, 67.13% bounce, 9 reach fees. NBA accreditation of B.Tech CSE, Biotechnology and MBA appears only here.",
            change="Certificate gallery → each accreditation gets a text headline + “Applies to: B.Tech CSE · B.Tech Biotechnology · MBA →” links. QS Asia 2026 quoted exactly (741–750 band). Verified claims only.",
            hyp="If each accreditation links to the programmes it covers, awards → programme progression rises, because credibility is attached to a decision.",
            measure="Primary: awards → programme rate · Secondary: leads · Guardrail: bounce (67.13%) · Scale if progression ≥ +25% relative",
            setup="50/50, 4 weeks. Effort S: content + links, no MREI data needed.",
            h=(0.52, 0.66, 0.40, 0.40, 0.28),
            mock="Control: certificate gallery\nVariant: text badges + programme links",
            page="/mriirs/about/awards-recognitions"),
       dict(id="PS-2", wave=2, name="Dated CSE placement proof beside the form", chips=["WAVE 2", "NEEDS MRIIRS DATA"],
            **{"from": "From pitch T4 (Fix 06) · MREI T8"},
            see="Salary cards (Microsoft 60 LPA, Amazon, Deloitte) exist in the code at 0×0 px. Testimonials: 2012–18 batches, looped 3×. “60 LPA” is unscoped. The CSE page never states its NBA accreditation.",
            change="DRDO / ISRO prose → “CSE batch 2025: [X]% placed · median ₹[Y] · [N] recruiters” + 3 named 2024–25 graduates (role, company, CTC) + NBA badge. Every figure sourced to the CRCDC report.",
            hyp="If placements are scoped, dated and named, the lead rate rises, because the payer’s ROI doubt is answered with evidence instead of a university-wide maximum.",
            measure="Primary: lead rate of sessions reaching placements · Secondary: form_start · Guardrail: engagement · Scale if ≥ +10% relative",
            setup="Blocked until MREI supplies CRCDC CSE data and confirms “60 LPA” and “95%”. Then 50/50, 4 weeks.",
            h=(0.64, 0.66, 0.52, 0.40, 0.40),
            mock="Control: prose + hidden cards\nVariant: CSE 2025 proof strip + NBA badge",
            page="B.Tech CSE page, placements section")],
      "PS-1 ships first; PS-2 waits for MREI’s placement data.",
      "GA4 explorations 13 Jun – 10 Sept 2026 · official MREI pages read 24 Sept 2026 · live DOM 23 Sept 2026 (MRIIRS pitch v4).",
      "Proof Seeker tests. No claim goes into a variant until MREI confirms it (see the proof inventory). " + RULE)

# ---------------------------------------------------------------- 5 University Shortlister
slide("University Shortlister",
      "University Shortlister: keep the benchmark, close the comparison leak",
      "“Is this university right for me?” · /mriirs and /mru · 76.1% engagement · 517 of 613 Apply-MRIIRS clicks",
      [dict(id="US-1", wave=1, name="Compare MRIIRS & MRU module", chips=["WAVE 1", "NEEDS MRIIRS DATA"],
            **{"from": "From MREI T12 · personalisation R2"},
            see="673 moves go from /mriirs to the homepage, 166 from /mru to /mriirs and 494 between the two fee pages. Nothing on the site supports the comparison.",
            change="Current hub → same hub + “MRIIRS or MRU? Compare in 30 seconds” below the UG / PG / Doctoral tiles: status · NAAC grade per institution · programmes · B.Tech CSE year-1 fee · campus · hostel, with “Help me choose — talk to a counsellor”.",
            hyp="If comparers can compare on the hub, leads rise and homepage loops fall, because the institution decision is made in one place instead of by hopping.",
            measure="Primary: leads · Secondary: cross-brand hops (166, 494) · Guardrail: Apply-MRIIRS clicks (517) must not fall",
            setup="50/50 on both roots, 4 weeks. Needs MRU B.Tech CSE fee and NAAC grade per institution confirmed.",
            h=(0.52, 0.78, 0.52, 0.40, 0.40),
            mock="Control: UG / PG / Doctoral hub\nVariant: hub + compare module",
            page="/mriirs · /mru institution roots"),
       dict(id="US-2", wave=2, name="Human routes instead of “Closing Soon”", chips=["WAVE 2", "SHIP-READY"],
            **{"from": "From pitch T7 (Fix 09) · MREI orchestration “Talk to a Counsellor”"},
            see="The only human route in the fold is a chat toast, “Admission Closing Soon”. City offices (Delhi, Guwahati, Indore, Kota, Lucknow, Varanasi, Patna, Hyderabad) sit in the footer.",
            change="Toast + generic sticky tabs → sticky “Talk to someone” with 3 tracked routes: book a campus visit (slot picker) · city-office counsellor, pre-selected by GA4 city · chat with a current student. Each issues a visit code.",
            hyp="If comparers can reach a person, campus visits rise, because the comparison is resolved in a conversation, not by hopping between copies of the site.",
            measure="Primary: visits booked and attended / 1K sessions · Secondary: visit → application · Guardrail: Apply clicks · Scale if ≥ +15% relative with SLA met",
            setup="50/50, 4 weeks, after US-1. Visit code captured at the counter closes the loop.",
            h=(0.52, 0.66, 0.52, 0.52, 0.28),
            mock="Control: chat toast + tabs\nVariant: “Talk to someone” sheet, 3 routes",
            page="/mriirs · /mru, sticky bottom bar")],
      "Add, don’t redesign: US-1 first, US-2 after.",
      "GA4 explorations 13 Jun – 10 Sept 2026 · GA4 event export by page type · live footer 23 Sept 2026 (MRIIRS pitch v4).",
      "University Shortlister tests. The strongest template on the site: every change is additive and guarded by Apply-MRIIRS clicks. " + RULE)

p.save(OUT)
print("\n".join(warn) or "no warnings")
